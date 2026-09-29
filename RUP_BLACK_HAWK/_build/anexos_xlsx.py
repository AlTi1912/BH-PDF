# Genera el libro de anexos editable a partir del modelo único.
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
M = json.load(open('model.json'))
E = M['ESTADOS']
wb = Workbook(); wb.remove(wb.active)
HDR = PatternFill('solid', fgColor='111111'); ALT = PatternFill('solid', fgColor='F6F1FE')
thin = Side(style='thin', color='D1D5DB')
def sheet(name, title, head, rows, widths):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=13, name='Arial')
    ws['A2'] = 'SGCD-BH · Proyecto académico RUP — Black Hawk Car Audio'; ws['A2'].font = Font(italic=True, size=9, color='6B7280', name='Arial')
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c.fill = HDR
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=5 + i, column=j, value=v)
            c.alignment = Alignment(vertical='top', wrap_text=True); c.font = Font(name='Arial', size=10)
            c.border = Border(bottom=thin)
            if i % 2: c.fill = ALT
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:{get_column_letter(len(head))}{4 + len(rows)}'
    return ws
cu = {c['id']: c for c in M['CU']}
sheet('RF', 'Catálogo de requisitos funcionales', ['Código', 'Nombre', 'Descripción', 'Prioridad', 'Actor', 'Caso de uso', 'Componente', 'Criterio de aceptación', 'Estado de implementación'],
      [[r['id'], r['n'], r['d'], r['p'], r['a'], r['cu'] + ' ' + cu[r['cu']]['n'], M['RF_COMP'][r['id']], r['ac'], E[r['e']]] for r in M['RF']], [10, 26, 50, 10, 20, 30, 12, 50, 24])
sheet('RNF', 'Catálogo de requisitos no funcionales', ['Código', 'Característica (ISO/IEC 25010)', 'Nombre', 'Descripción', 'Prioridad', 'Criterio de aceptación', 'Estado'],
      [[r['id'], r['c'], r['n'], r['d'], r['p'], r['ac'], E[r['e']]] for r in M['RNF']], [10, 18, 26, 55, 10, 50, 24])
sheet('Casos de uso', 'Inventario de casos de uso', ['Código', 'Nombre', 'Actor(es)', 'Módulo', 'Relaciones', 'Estado'],
      [[c['id'], c['n'], c['a'], next(m['n'] for m in M['MODULOS'] if m['id'] == c['m']), c.get('rel', ''), E[c['e']]] for c in M['CU']], [9, 38, 36, 34, 18, 26])
comp = {c['id']: c['n'] for c in M['COMPONENTES']}; nb = {n['id']: n['t'] for n in M['NECESIDADES']}
sheet('Trazabilidad', 'Matriz de trazabilidad: necesidad → requisito → caso de uso → componente → prueba', ['Necesidad', 'Descripción', 'Requisitos', 'Casos de uso', 'Componentes', 'Casos de prueba'],
      [[t['nb'], nb[t['nb']], ', '.join(t['rf']), ', '.join(t['cu']), ', '.join(f"{c} {comp[c]}" for c in t['comp']), ', '.join(t['cp'])] for t in M['TRAZA']], [10, 40, 30, 28, 44, 30])
sheet('Plan de pruebas', 'Casos de prueba (diseñados; pendientes de ejecución)', ['Código', 'Requisito', 'Tipo', 'Caso de prueba', 'Resultado esperado', 'Estado', 'Resultado obtenido'],
      [[c['id'], c['rf'], c['t'], c['n'], c['r'], 'Diseñado — no ejecutado', ''] for c in M['CP']], [9, 10, 12, 42, 44, 22, 22])
def nivel(s): return 'Crítico' if s >= 16 else 'Alto' if s >= 12 else 'Medio' if s >= 6 else 'Bajo'
sheet('Riesgos', 'Matriz de riesgos', ['Código', 'Riesgo', 'Categoría', 'Probabilidad (1-5)', 'Impacto (1-5)', 'P × I', 'Nivel', 'Estrategia de mitigación', 'Responsable'],
      [[r['id'], r['n'], r['cat'], r['p'], r['i'], r['p'] * r['i'], nivel(r['p'] * r['i']), r['m'], r['resp']] for r in M['RIESGOS']], [9, 44, 12, 12, 12, 8, 10, 60, 20])
sheet('Cronograma', 'Cronograma por iteraciones (semanas académicas)', ['Iteración', 'Fase', 'Semanas', 'Objetivo', 'Casos de uso', 'Entregable'],
      [[i['id'], i['f'], f"S{i['s'][0]}–S{i['s'][1]}", i['obj'], i['cu'], i['ent']] for i in M['ITERACIONES']], [10, 14, 10, 55, 36, 30])
sheet('EDT', 'Estructura de desglose del trabajo (EDT/WBS)', ['Código', 'Paquete de trabajo'], M['EDT'], [10, 50])
tot = sum(e['h'] for e in M['ESFUERZO'])
rows = [[e['f'], e['h'], round(100 * e['h'] / tot, 1), M['TARIFA_REF'], e['h'] * M['TARIFA_REF']] for e in M['ESFUERZO']]
rows.append(['Total', tot, 100.0, M['TARIFA_REF'], tot * M['TARIFA_REF']])
ws = sheet('Esfuerzo', 'Estimación de esfuerzo y presupuesto referencial (cifras ESTIMADAS, supuesto académico)', ['Fase', 'Horas-persona', '% del esfuerzo', 'Tarifa hipotética (S/ por hora)', 'Costo referencial (S/)'], rows, [16, 16, 16, 26, 22])
ws.cell(row=6 + len(rows), column=1, value='Nota: la tarifa de S/ 30 por hora es un valor hipotético para el ejercicio académico. No representa un costo real de Black Hawk ni una cotización.').font = Font(italic=True, size=9, name='Arial')
drows = []
for t, cols in M['DICCIONARIO'].items():
    for c in cols: drows.append([t] + c)
sheet('Diccionario de datos', 'Diccionario de datos (modelo lógico)', ['Tabla', 'Columna', 'Tipo', 'Restricción', 'Descripción'], drows, [22, 22, 14, 36, 40])
sheet('Mapeo WordPress', 'Correspondencia del modelo lógico con el almacenamiento físico en WordPress', ['Entidad', 'Almacenamiento físico', 'Estado'], [[a, b, E[c]] for a, b, c in M['MAPEO']], [26, 55, 26])
sheet('Diagramas', 'Inventario de diagramas', ['Código', 'Archivo (PNG/SVG/PUML)', 'Título', 'Tipo UML', 'Elementos relacionados'], M['DIAGRAMAS'], [8, 30, 50, 24, 36])
sheet('Glosario', 'Glosario técnico', ['Término', 'Definición'], M['GLOSARIO'], [28, 90])
sheet('Stakeholders', 'Identificación de stakeholders', ['Stakeholder', 'Rol', 'Interés', 'Influencia', 'Interés en el proyecto'], [[s['n'], s['r'], s['i'], s['inf'], s['int']] for s in M['STAKEHOLDERS']], [30, 18, 60, 12, 14])
wb.save('../ANEXOS/BLACK_HAWK_RUP_ANEXOS.xlsx'); print('ok')
