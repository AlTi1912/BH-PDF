# Genera ANEXOS/BLACK_HAWK_RUP_ANEXOS_V3.xlsx: parte del libro V2 y solo corrige lo que cambia con el enfoque V3
# (sitio original → web renovada → ampliaciones). Requiere v3data.json (ver README).
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
M = json.load(open('model.json'))
V = json.load(open('v3data.json'))
ORIG = json.load(open('v3orig.json'))
wb = load_workbook('../ANEXOS/BLACK_HAWK_RUP_ANEXOS_V2.xlsx')
HDR = PatternFill('solid', fgColor='111111'); ALT = PatternFill('solid', fgColor='F6F1FE')
thin = Side(style='thin', color='D1D5DB')
SUB = 'Proyecto académico RUP — Black Hawk Car Audio · V3 (sitio original → web renovada → ampliaciones)'

def sheet(name, title, head, rows, widths, pos=None):
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=13, name='Arial')
    ws['A2'] = SUB; ws['A2'].font = Font(italic=True, size=9, color='6B7280', name='Arial')
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c.fill = HDR
        c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=5 + i, column=j, value=v)
            c.alignment = Alignment(vertical='top', wrap_text=True); c.font = Font(name='Arial', size=10)
            c.border = Border(bottom=thin)
            if i % 2: c.fill = ALT
    for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:{get_column_letter(len(head))}{4 + len(rows)}'
    if pos is not None: wb.move_sheet(ws, offset=pos - wb.sheetnames.index(name))
    return ws

def note(ws, texts):
    r = ws.max_row + 2
    for t in texts:
        ws.cell(row=r, column=1, value=t).font = Font(name='Arial', size=9, italic=True); r += 1

def set_col(ws, col, head, values, width):
    c0 = ws.cell(row=4, column=col, value=head); c0.font = Font(bold=True, color='FFFFFF', name='Arial', size=10); c0.fill = HDR
    c0.alignment = Alignment(vertical='center', wrap_text=True)
    for r in range(5, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if code is None: continue
        c = ws.cell(row=r, column=col, value=values(code)); c.font = Font(name='Arial', size=10)
        c.alignment = Alignment(vertical='top', wrap_text=True); c.border = Border(bottom=thin)
        if (r - 5) % 2: c.fill = ALT
    ws.column_dimensions[get_column_letter(col)].width = width

T = V['TIPOS']
F = V['FECHA']

# 1. Sitio original: punto de partida, hallazgos clasificados
rows = [['Existente', t, '—', f'Sitio original, {F}'] for t in V['ORIGINAL_TIENE']]
rows += [[T[h['tipo']], f"{h['n']}. {h['t']}", h['d'], h['ev']] for h in V['HALLAZGOS']]
rows += [[T[t], h, '—', ev] for h, t, ev in V['HALLAZGOS_EXTRA']]
ws = sheet('Sitio original', 'Sitio original (blackhawkcaraudio.com): lo que existe y los hallazgos, clasificados', ['Tipo', 'Elemento o hallazgo', 'Descripción', 'Evidencia'], rows, [22, 50, 70, 30], pos=1)
note(ws, ['Tipos: Existente · Inexistente · Existe, poco visible · Recorrido mejorable. Los números 1–6 coinciden con los marcadores de la diapositiva 4 de la presentación V3.',
          'Las capturas del «antes» son del sitio original; nunca se usaron capturas de la web renovada para representar el estado anterior.'])

# 2. Antes y después
rows = [[a['n'], a['a'], a['nec'], a['sol'], a['ben'], f"{a['antes'].split('/')[-1]} → {a['despues'].split('/')[-1]}"] for a in V['ANTES_DESPUES']]
ws = sheet('Antes y después', 'Comparaciones entre pantallas equivalentes', ['Pantalla', 'Antes (sitio original)', 'Necesidad', 'Modificación (web renovada)', 'Beneficio esperado', 'Capturas (_build/img/)'], rows, [18, 40, 36, 40, 34, 30], pos=2)
rows_m = [[a, b, c] for a, b, c in V['MEDICIONES']]
r0 = ws.max_row + 2
ws.cell(row=r0, column=1, value='Mediciones comparables (Lighthouse, medianas, 23-09-2026)').font = Font(bold=True, name='Arial', size=10)
for i, (a, b, c) in enumerate(rows_m, 1):
    for j, v in enumerate([a, b, c], 1): ws.cell(row=r0 + i, column=j, value=v).font = Font(name='Arial', size=10)
note(ws, ['«Beneficio esperado» no es un resultado medido: no se dispone de datos de uso, Analytics ni ventas. Un clic en «Cotizar» abre WhatsApp; no equivale a una venta.'])

# 3. Datos verificados (reemplaza la hoja V2: separa sitio original y web renovada)
datos = [
    ['Productos y categorías del sitio original', '110 productos · 11 categorías', 'Sitemap de blackhawkcaraudio.com', F, 'Diapositiva 2'],
    ['Página /shop/ del sitio original', 'Título «Tienda» sin productos', 'Captura del sitio original', F, 'Diapositivas 4 y 21'],
    ['Enlaces wa.me o tel: en home, ficha y categoría del sitio original', '0', 'Revisión de enlaces del sitio original', F, 'Diapositiva 4'],
    ['WhatsApp en el sitio original', 'Solo en la página «Ventas al mayor»', 'Revisión de enlaces del sitio original', F, 'Diapositiva 4'],
    ['Productos y categorías de la web renovada', '118 productos · 12 categorías (se agrega «Cargadores»)', 'Sitemaps de la web renovada', F, 'Diapositiva 21'],
    ['Máximo de productos en el comparador', '3', 'Código del tema (navigation.js)', F, 'Diapositivas 11 y 24'],
    ['Vigencia de la selección del comparador', '24 h, en el navegador', 'Código del tema (navigation.js)', F, 'Diapositiva 24'],
    ['Mensaje de WhatsApp desde la ficha', '«Hola Black Hawk, vengo de la web y quiero cotizar el modelo …» (sin URL)', 'Enlaces wa.me de las fichas', F, 'Diapositivas 17 y 23'],
    ['Directorio de distribuidores', 'No existe: «Dónde comprar» deriva a WhatsApp (ampliación propuesta)', 'Página /distribuidores/', F, 'Diapositivas 23 y 24'],
] + [[a, f'{b} → {c}', 'Lighthouse, medianas (evidence/data/lighthouse-medianas.json)', '23-09-2026', 'Diapositivas 14 y 25'] for a, b, c in V['MEDICIONES']] + [
    ['Estado de las fases', 'Inicio y Elaboración realizadas; Construcción en curso; Transición planificada', 'Declaración del proyecto', F, 'Diapositivas 9–13'],
]
ws = sheet('Datos verificados', 'Cifras y afirmaciones usadas en la V3, con su fuente', ['Dato', 'Valor', 'Fuente', 'Fecha', 'Dónde se usa (presentación V3)'], datos, [40, 44, 40, 12, 26], pos=3)
note(ws, ['Las estimaciones de la V1 (14 semanas, 560 h, S/ 16 800 con tarifa hipotética) no se usan en la V3; se conservan en la hoja «Esfuerzo» como supuesto académico.'])

# 4. RF y RNF: estado en el sitio original y en la web renovada (sustituye «En la presentación V2»)
V3E = {'EXI': 'Desarrollado', 'BASE': 'Base WordPress/WooCommerce', 'PAR': 'Parcial', 'PEN': 'Pendiente de validación', 'PRO': 'Ampliación propuesta'}
DET = {
    'RF-004': 'Parcial: especificaciones en las 118 fichas (112 en lista, 6 en párrafo), aún sin datos estructurados',
    'RF-015': 'Parcial: página Mayoristas con formulario y consentimiento; validación del RUC por verificar',
    'RF-019': 'Parcial: especificaciones como texto en la ficha, sin atributo-valor-unidad ni fuente',
    'RF-010': 'Pendiente de validación con Black Hawk',
    'RF-027': 'Pendiente: por verificar con el hosting',
}
rfe = {r['id']: r['e'] for r in M['RF']}
en_ppt_rf = {c.strip() for _, _, cs in V['REQ_F'] for c in cs.split(',')}
en_ppt_rnf = {c for _, _, c in V['REQ_NF']}
ws = wb['RF']
set_col(ws, 10, 'Sitio original', lambda c: ORIG.get(c, '—'), 34)
set_col(ws, 11, 'Web renovada / proyecto (V3)', lambda c: DET.get(c, V3E[rfe[c]]), 30)
set_col(ws, 12, 'En la presentación V3', lambda c: 'Sí (diap. 14)' if c in en_ppt_rf else '—', 14)
ws.auto_filter.ref = f'A4:L{ws.max_row}'
ws = wb['RNF']
set_col(ws, 8, 'En la presentación V3', lambda c: 'Sí (diap. 14)' if c in en_ppt_rnf else '—', 14)
ws.auto_filter.ref = f'A4:H{ws.max_row}'

# 5. Iteraciones reales del proyecto (la hoja «Cronograma» conserva el plan V1)
disc = [d for d, _ in V['MATRIZ']]
rows = []
for k, it in enumerate(V['ITER']):
    ent = '; '.join(f'{d}: {cells[k]}' for d, cells in V['MATRIZ'] if cells[k])
    rows.append([it['id'], it['f'], it['e'], ent])
ws = sheet('Iteraciones V3', 'Iteraciones del proyecto y entregables por disciplina', ['Iteración', 'Fase', 'Estado', 'Entregables principales'], rows, [10, 14, 14, 110], pos=4)
note(ws, ['El orden y el contenido de las iteraciones reconstruyen el proceso a partir de la evidencia del rediseño; ajústalos si el proceso real fue distinto.',
          'La hoja «Cronograma» conserva el plan de la V1, que asignaba el registro de consultas y los distribuidores a Construcción; en la V3 esas funciones son ampliaciones propuestas.'])
ws = wb['Cronograma']
ws['A2'] = 'Plan de la V1 (supuesto académico). Sustituido en la V3 por la hoja «Iteraciones V3».'

# 6. Inventario de diagramas: agrega los de la V3
ws = wb['Diagramas']
v3d = [
    ['V3-01', 'V3_01_casos_de_uso', 'Casos de uso de la web renovada (sitio público y administración; propuestos marcados)', 'Casos de uso', 'Presentación V3 diap. 15; monografía V3'],
    ['V3-02', 'V3_02_actividades', 'Recorrido comercial: cliente, web renovada y área comercial', 'Actividades', 'Presentación V3 diap. 16; monografía V3'],
    ['V3-03', 'V3_03_secuencia', 'Cotizar por WhatsApp desde la ficha', 'Secuencia', 'Presentación V3 diap. 17; monografía V3'],
    ['V3-04', 'V3_04_clases', 'Clases del catálogo (Distribuidor y EventoConsulta propuestos)', 'Clases', 'Presentación V3 diap. 18; monografía V3'],
    ['V3-05', 'V3_05_componentes', 'Componentes: tema hijo rozer-child sobre WooCommerce y WordPress', 'Componentes', 'Presentación V3 diap. 19; monografía V3'],
    ['D-20', 'D20_despliegue', 'Despliegue (respaldo)', 'Despliegue', 'Presentación V3 diap. 28 (respaldo)'],
]
start = ws.max_row + 1
for i, r in enumerate(v3d):
    for j, v in enumerate(r, 1):
        c = ws.cell(row=start + i, column=j, value=v); c.font = Font(name='Arial', size=10); c.alignment = Alignment(vertical='top', wrap_text=True); c.border = Border(bottom=thin)
ws.auto_filter.ref = f'A4:E{ws.max_row}'

# 7. LÉEME
niveles = [
    ['Nivel 1 · Presentación V3', 'Sitio original', 'Punto de partida: lo que ya existía y los hallazgos clasificados (inexistente, poco visible, recorrido mejorable).'],
    ['Nivel 1 · Presentación V3', 'Antes y después', 'Comparaciones entre pantallas equivalentes y mediciones de peso de página.'],
    ['Nivel 1 · Presentación V3', 'Datos verificados', 'Cifras de la presentación, separadas por sitio original y web renovada, con su fuente.'],
    ['Nivel 2 · Monografía V3', 'Iteraciones V3', 'Iteraciones reales del proyecto, estado y entregables por disciplina.'],
    ['Nivel 2 · Monografía V3', 'RF (columnas «Sitio original» y «Web renovada / proyecto»)', 'Estado de cada requisito en el punto de partida y en la solución.'],
    ['Nivel 2 · Monografía V3', 'Productos de referencia, Casos de uso, Stakeholders, Riesgos', 'Se conservan de la V2.'],
    ['Nivel 3 · Anexos técnicos', 'RF, RNF, Trazabilidad, Plan de pruebas', 'Catálogo completo con criterios de aceptación; 32 casos de prueba diseñados, no ejecutados.'],
    ['Nivel 3 · Anexos técnicos', 'Cronograma, EDT, Esfuerzo', 'Plan y estimaciones de la V1 (supuesto académico).'],
    ['Nivel 3 · Anexos técnicos', 'Diccionario de datos, Mapeo WordPress, Diagramas, Glosario', 'Modelo de datos (incluye las ampliaciones propuestas) e inventario de diagramas, con los V3-01 a V3-05.'],
]
ws = sheet('LÉEME', 'Anexos del proyecto · versión 3 · índice por niveles', ['Nivel', 'Hoja', 'Contenido'], niveles, [26, 50, 90], pos=0)
note(ws, ['Tres partes: A) sitio original (punto de partida, existente) · B) web renovada (desarrollada, en entorno de prueba) · C) ampliaciones (propuestas).',
          'Estados de requisitos en la V3: Desarrollado · Base WordPress/WooCommerce · Parcial · Pendiente de validación · Ampliación propuesta.',
          f'Datos del sitio original y de la web renovada verificados el {F}; mediciones Lighthouse del 23-09-2026. Sin datos de Analytics ni de ventas.',
          'Las pruebas formales están diseñadas, no ejecutadas; las horas y el presupuesto son estimaciones académicas.'])
wb.active = 0
wb.save('../ANEXOS/BLACK_HAWK_RUP_ANEXOS_V3.xlsx')
print(wb.sheetnames)
